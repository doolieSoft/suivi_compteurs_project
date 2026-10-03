"""Export de toutes les données dans un classeur Excel.

Destiné à être conservé hors de l'application (Google Drive, clé USB). La
présentation reprend celle de ``compteur.xlsx`` : une feuille par énergie, un
bloc par compteur, l'index relevé suivi des écarts calculés par formule — le
fichier reste donc vivant si on y ajoute une ligne à la main. Suivent une
synthèse annuelle et les données de référence (maisons, compteurs, événements,
tarifs, degrés-jours).

Le classeur se réimporte (``import_excel`` le reconnaît) : chaque bloc de
relevés commence par le *code* du compteur, qui renvoie à sa ligne de la
feuille « Compteurs ». Seules les colonnes saisies sont relues — date, index,
relevé annuel, source, commentaire — ; les colonnes calculées et la feuille
« Par année » sont ignorées, puisque l'application les recalcule.
"""
from __future__ import annotations

import datetime as dt
import io
from collections.abc import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from django.utils.text import slugify

from suivi.models import Compteur, DegreJour, Energie, Evenement, Maison, Tarif
from suivi.services import consommation as cs
from suivi.services import degres_jours as djs

FORMAT_DATE = "DD/MM/YYYY"
GRAS = Font(bold=True)
TITRE = Font(bold=True, size=12)

FEUILLES_ENERGIE = (
    (Energie.EAU, "Eau"),
    (Energie.GAZ, "Gaz"),
    (Energie.ELECTRICITE, "Électricité"),
    (Energie.MAZOUT, "Mazout"),
)

ENTETES_RELEVES = (
    "Date", "Index", "Écart", "Jours", "Par jour", "Sur un an",
    "Relevé annuel", "Source", "Commentaire",
)
LARGEURS_RELEVES = (12, 13, 11, 7, 10, 11, 13, 18, 40)

# Noms et en-têtes relus à l'import : les changer impose d'adapter celui-ci.
FEUILLE_MAISONS = "Maisons"
FEUILLE_COMPTEURS = "Compteurs"
FEUILLE_EVENEMENTS = "Événements"
FEUILLE_TARIFS = "Tarifs"
FEUILLE_DJ = "Degrés-jours"
ENTETE_CODE = "Code"

ENTETES_MAISONS = (
    ENTETE_CODE, "Nom", "Adresse", "Façades", "Surface chauffée (m²)",
    "Occupée depuis", "Occupée jusqu'au", "Station météo", "Latitude",
    "Longitude", "Base des degrés-jours", "Notes",
)
ENTETES_COMPTEURS = (
    ENTETE_CODE, "Maison", "Énergie", "Plage", "Libellé", "Numéro", "Unité",
    "Coefficient kWh", "Posé le", "Déposé le", "Remplace",
)
ENTETES_EVENEMENTS = ("Maison", "Date", "Énergie", "Libellé", "Description")
ENTETES_TARIFS = (
    "Maison", "Énergie", "Fournisseur", "Début", "Fin", "Prix unitaire (€)",
    "Abonnement (€/mois)", "Notes",
)
ENTETES_DJ = ("Station", "Date", "Température moyenne (°C)", "Degrés-jours")


def codes_compteurs(compteurs: Iterable[Compteur]) -> dict[int, str]:
    """Code lisible et unique de chaque compteur, par identifiant."""
    codes: dict[int, str] = {}
    pris: set[str] = set()
    for c in compteurs:
        base = "-".join(
            part
            for part in (
                c.maison.slug,
                slugify(c.get_energie_display()),
                slugify(c.get_plage_display()),
                slugify(c.libelle),
            )
            if part
        )
        code, n = base, 2
        while code in pris:
            code, n = f"{base}-{n}", n + 1
        pris.add(code)
        codes[c.pk] = code
    return codes


def _feuille_releves(
    feuille: Worksheet, compteurs: Sequence[Compteur], codes: dict[int, str]
) -> None:
    """Un bloc par compteur, empilés comme dans le classeur d'origine."""
    for rang, largeur in enumerate(LARGEURS_RELEVES, start=1):
        feuille.column_dimensions[get_column_letter(rang)].width = largeur

    ligne = 1
    for compteur in compteurs:
        releves = list(compteur.releves.order_by("date"))
        if not releves:
            continue

        # Le code, en tête de bloc, est ce que l'import reconnaît.
        feuille.cell(ligne, 1, codes[compteur.pk]).font = TITRE
        details = [str(compteur), compteur.unite]
        if compteur.numero:
            details.append(f"n° {compteur.numero}")
        if compteur.date_depose:
            details.append(f"déposé le {compteur.date_depose:%d/%m/%Y}")
        feuille.cell(ligne, 3, " · ".join(details))
        ligne += 1

        for colonne, entete in enumerate(ENTETES_RELEVES, start=1):
            feuille.cell(ligne, colonne, entete).font = GRAS
        ligne += 1

        premiere = ligne
        for releve in releves:
            feuille.cell(ligne, 1, releve.date).number_format = FORMAT_DATE
            feuille.cell(ligne, 2, float(releve.index)).number_format = "0.000"
            if ligne > premiere:
                # Formules plutôt que valeurs : comme dans compteur.xlsx, une
                # ligne ajoutée à la main se calcule toute seule.
                p = ligne - 1
                feuille.cell(ligne, 3, f"=B{ligne}-B{p}").number_format = "0.000"
                feuille.cell(ligne, 4, f"=A{ligne}-A{p}").number_format = "0"
                feuille.cell(ligne, 5, f'=IF(D{ligne}>0,C{ligne}/D{ligne},"")').number_format = "0.000"
                feuille.cell(ligne, 6, f'=IF(D{ligne}>0,E{ligne}*365,"")').number_format = "0"
            feuille.cell(ligne, 7, "oui" if releve.annuel else None)
            feuille.cell(ligne, 8, releve.get_source_display())
            feuille.cell(ligne, 9, releve.commentaire or None)
            ligne += 1

        ligne += 1  # ligne vide entre deux compteurs

    feuille.freeze_panes = "A1"


def _feuille_tableau(
    feuille: Worksheet,
    entetes: Sequence[str],
    lignes: Iterable[Sequence],
    *,
    formats: dict[int, str] | None = None,
    largeurs: Sequence[int] | None = None,
) -> None:
    """Tableau simple : en-tête en gras et figé."""
    feuille.append(list(entetes))
    for cellule in feuille[1]:
        cellule.font = GRAS
    feuille.freeze_panes = "A2"
    for ligne in lignes:
        feuille.append(list(ligne))
    for rang, format_colonne in (formats or {}).items():
        for (cellule,) in feuille.iter_rows(min_row=2, min_col=rang, max_col=rang):
            cellule.number_format = format_colonne
    for rang, largeur in enumerate(largeurs or [max(12, len(e) + 2) for e in entetes], start=1):
        feuille.column_dimensions[get_column_letter(rang)].width = largeur


def classeur() -> bytes:
    """Renvoie le contenu du fichier .xlsx."""
    resultat = Workbook()
    resultat.remove(resultat.active)

    # Maison actuelle d'abord, comme on la consulte le plus souvent.
    ordre_maisons = {
        m.pk: rang
        for rang, m in enumerate(
            sorted(Maison.objects.all(), key=lambda m: (m.date_sortie is not None, m.nom))
        )
    }
    compteurs = sorted(
        Compteur.objects.select_related("maison"),
        key=lambda c: (
            ordre_maisons.get(c.maison_id, 99),
            c.plage,
            c.date_pose or dt.date.min,
        ),
    )
    codes = codes_compteurs(compteurs)
    for energie, titre in FEUILLES_ENERGIE:
        _feuille_releves(
            resultat.create_sheet(titre),
            [c for c in compteurs if c.energie == energie],
            codes,
        )

    synthese = []
    for maison in sorted(Maison.objects.select_related("station"), key=lambda m: ordre_maisons[m.pk]):
        normales = djs.dj_normaux_par_jour_calendaire(maison.station)
        for ligne in cs.lignes_de(maison):
            for annee in cs.comparer_annees(ligne, normales=normales):
                synthese.append(
                    [
                        maison.nom,
                        ligne.libelle_energie,
                        annee.annee,
                        round(annee.consommation, 3),
                        ligne.unite,
                        annee.jours_couverts,
                        "oui" if annee.complete else "non",
                        round(annee.dj_reel, 1),
                        None
                        if annee.consommation_normalisee is None
                        else round(annee.consommation_normalisee, 3),
                    ]
                )
    _feuille_tableau(
        resultat.create_sheet("Par année"),
        ["Maison", "Énergie", "Année", "Consommation", "Unité", "Jours couverts",
         "Année complète", "Degrés-jours", "À climat normal"],
        synthese,
        formats={4: "0.0", 8: "0", 9: "0.0"},
        largeurs=[14, 28, 8, 14, 7, 15, 15, 13, 16],
    )

    maisons = sorted(Maison.objects.select_related("station"), key=lambda m: ordre_maisons[m.pk])
    _feuille_tableau(
        resultat.create_sheet(FEUILLE_MAISONS),
        ENTETES_MAISONS,
        (
            [
                m.slug,
                m.nom,
                m.adresse,
                m.nb_facades,
                m.surface_m2,
                m.date_entree,
                m.date_sortie,
                m.station.nom if m.station else None,
                float(m.station.latitude) if m.station else None,
                float(m.station.longitude) if m.station else None,
                float(m.station.base_dj) if m.station else None,
                m.notes,
            ]
            for m in maisons
        ),
        formats={6: FORMAT_DATE, 7: FORMAT_DATE, 9: "0.00000", 10: "0.00000"},
    )

    _feuille_tableau(
        resultat.create_sheet(FEUILLE_COMPTEURS),
        ENTETES_COMPTEURS,
        (
            [
                codes[c.pk],
                c.maison.slug,
                c.get_energie_display(),
                c.get_plage_display(),
                c.libelle,
                c.numero,
                c.unite,
                float(c.coef_kwh),
                c.date_pose,
                c.date_depose,
                codes.get(c.remplace_id) if c.remplace_id else None,
            ]
            for c in compteurs
        ),
        formats={9: FORMAT_DATE, 10: FORMAT_DATE},
        largeurs=[34, 14, 12, 15, 16, 14, 7, 16, 12, 12, 34],
    )

    _feuille_tableau(
        resultat.create_sheet(FEUILLE_EVENEMENTS),
        ENTETES_EVENEMENTS,
        (
            [e.maison.slug, e.date, e.get_energie_display(), e.libelle, e.description]
            for e in Evenement.objects.select_related("maison").order_by("date")
        ),
        formats={2: FORMAT_DATE},
        largeurs=[14, 12, 12, 40, 60],
    )

    _feuille_tableau(
        resultat.create_sheet(FEUILLE_TARIFS),
        ENTETES_TARIFS,
        (
            [
                t.maison.slug if t.maison else None,
                t.get_energie_display(),
                t.fournisseur,
                t.date_debut,
                t.date_fin,
                float(t.prix_unitaire),
                float(t.abonnement_mensuel),
                t.notes,
            ]
            for t in Tarif.objects.select_related("maison").order_by("energie", "date_debut")
        ),
        formats={4: FORMAT_DATE, 5: FORMAT_DATE, 6: "0.00000", 7: "0.00"},
    )

    _feuille_tableau(
        resultat.create_sheet(FEUILLE_DJ),
        ENTETES_DJ,
        (
            [nom, jour, None if temperature is None else float(temperature), float(valeur)]
            for nom, jour, temperature, valeur in DegreJour.objects.order_by(
                "station__nom", "date"
            ).values_list("station__nom", "date", "temperature_moyenne", "dj")
        ),
        formats={2: FORMAT_DATE, 3: "0.00", 4: "0.000"},
        largeurs=[12, 12, 26, 13],
    )

    tampon = io.BytesIO()
    resultat.save(tampon)
    return tampon.getvalue()
