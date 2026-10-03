"""Import du classeur ``compteur.xlsx``.

Le classeur mélange, dans une même feuille, des blocs de nature différente :
index annuels de l'ancienne maison, relevés détaillés de la maison actuelle,
et des calculs intermédiaires. La constante ``BLOCS`` décrit donc explicitement
où lire quoi, plutôt que de deviner.

Deux automatismes évitent d'avoir à décrire les cas particuliers :

* un index qui **recule** signale un remplacement de compteur : la série est
  découpée en deux ``Compteur`` distincts, ce qui empêche de calculer un écart
  absurde à la charnière ;
* les cellules texte voisines d'un relevé (« index envoyé », « changement de
  douche »…) sont reprises en commentaire.

Import idempotent : relancer la commande met à jour les relevés existants
plutôt que de les dupliquer.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.text import slugify

from suivi.models import (
    Compteur,
    DegreJour,
    Energie,
    Evenement,
    Maison,
    Plage,
    Releve,
    SourceReleve,
    StationMeteo,
    Tarif,
)
from suivi.services import export_excel as ex

# Liège : les deux maisons sont dans la même région climatique.
STATION_DEFAUT = {
    "nom": "Liège",
    "latitude": Decimal("50.63730"),
    "longitude": Decimal("5.57950"),
}

MAISONS = {
    "collectivite": {
        "nom": "Collectivité",
        "nb_facades": 2,
        "date_entree": None,
        "date_sortie": dt.date(2022, 10, 31),
        "notes": "Ancienne maison. Index relevés annuellement jusqu'au déménagement.",
    },
    "liserons": {
        "nom": "Liserons",
        "nb_facades": 4,
        "date_entree": dt.date(2022, 4, 1),
        "date_sortie": None,
        "notes": "Maison actuelle. Acquise en novembre 2021, occupée depuis avril 2022.",
    },
}


@dataclass(frozen=True)
class Bloc:
    """Emplacement d'une série de relevés dans le classeur."""

    feuille: str
    maison: str
    energie: str
    ligne_debut: int
    ligne_fin: int
    col_date: str
    colonnes: dict[str, str]  # plage tarifaire -> lettre de colonne
    col_commentaires: tuple[str, ...] = ()
    unite: str = "m³"


BLOCS: tuple[Bloc, ...] = (
    # --- Ancienne maison : un index par an -------------------------------
    Bloc("Eau", "collectivite", Energie.EAU, 3, 12, "A", {Plage.UNIQUE: "B"}, ("C",)),
    Bloc("Gaz", "collectivite", Energie.GAZ, 4, 12, "A", {Plage.UNIQUE: "B"}, ("C",)),
    Bloc(
        "Electricité",
        "collectivite",
        Energie.ELECTRICITE,
        2,
        10,
        "A",
        {Plage.HAUT: "B", Plage.BAS: "C"},
        (),
        unite="kWh",
    ),
    # --- Maison actuelle : relevés détaillés ------------------------------
    Bloc(
        "Eau",
        "liserons",
        Energie.EAU,
        25,
        68,
        "A",
        {Plage.UNIQUE: "B"},
        ("C", "D", "F"),
    ),
    Bloc("Gaz", "liserons", Energie.GAZ, 29, 74, "A", {Plage.UNIQUE: "B"}, ("D", "G")),
    Bloc(
        "Electricité",
        "liserons",
        Energie.ELECTRICITE,
        28,
        75,
        "A",
        {Plage.UNIQUE: "B"},
        ("D", "E"),
        unite="kWh",
    ),
)

# Les textes ci-dessous marquent des travaux, pas de simples annotations de relevé.
MOTS_CLES_EVENEMENT = (
    "changement",
    "isolation",
    "remplacement",
    "travaux",
)

# Feuille « Prix par année » : ligne du millésime -> ligne des montants.
BLOCS_TARIFS = (
    (1, 3, None),
    (5, 7, 8),
    (10, 12, 13),
    (15, 17, 18),
)
COLONNES_TARIFS = {"A": Energie.ELECTRICITE, "B": Energie.EAU, "C": Energie.GAZ}


def _vers_date(valeur) -> dt.date | None:
    """Convertit une cellule en date, en rejetant tout ce qui n'en est pas une."""
    if isinstance(valeur, dt.datetime):
        return valeur.date()
    if isinstance(valeur, dt.date):
        return valeur
    if isinstance(valeur, (int, float)) and float(valeur).is_integer():
        entier = int(valeur)
        # Un millésime : l'index annuel est daté du 31 décembre.
        if 1990 <= entier <= 2100:
            return dt.date(entier, 12, 31)
        # Un numéro de série Excel plausible (1990-2100).
        if 32874 <= entier <= 73415:
            return dt.date(1899, 12, 30) + dt.timedelta(days=entier)
    return None


def _vers_decimal(valeur, decimales: int = 3) -> Decimal | None:
    if valeur is None or isinstance(valeur, (str, bool)):
        return None
    try:
        nombre = Decimal(str(valeur))
    except (InvalidOperation, ValueError):
        return None
    if nombre < 0:
        return None
    return nombre.quantize(Decimal(1).scaleb(-decimales))


def _vers_jour(valeur) -> dt.date | None:
    """Date d'une cellule de l'export : jamais un millésime, contrairement à ``_vers_date``."""
    if isinstance(valeur, dt.datetime):
        return valeur.date()
    if isinstance(valeur, dt.date):
        return valeur
    if isinstance(valeur, str):
        for format_ in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return dt.datetime.strptime(valeur.strip(), format_).date()
            except ValueError:
                pass
    return None


def _vers_entier(valeur) -> int | None:
    if isinstance(valeur, (int, float)) and not isinstance(valeur, bool):
        return int(valeur)
    return None


def _vers_coordonnee(valeur) -> Decimal:
    nombre = Decimal(str(valeur if valeur is not None else 0))
    return nombre.quantize(Decimal("0.00001"))


def _texte(valeur) -> str:
    return "" if valeur is None else str(valeur).strip()


def _choix(choix, valeur) -> str | None:
    """Code d'un choix à partir de son libellé (« Gaz ») ou de son code (« GAZ »)."""
    texte = _texte(valeur)
    for code, libelle in choix.choices:
        if texte in (code, libelle):
            return code
    return None


def _est_un_export(classeur) -> bool:
    """Le classeur vient-il de l'export de l'application ?"""
    if ex.FEUILLE_COMPTEURS not in classeur.sheetnames:
        return False
    premiere = next(classeur[ex.FEUILLE_COMPTEURS].iter_rows(max_row=1, values_only=True), ())
    return bool(premiere) and premiere[0] == ex.ENTETE_CODE


def _lignes(feuille, entetes):
    """Lignes d'un tableau de l'export, en dictionnaires indexés par en-tête.

    Les colonnes sont retrouvées par leur titre : une colonne déplacée ou
    ajoutée à la main dans le classeur ne fausse donc pas la lecture.
    """
    lignes = feuille.iter_rows(values_only=True)
    titres = [_texte(t) for t in next(lignes, ())]
    positions = {titre: titres.index(titre) for titre in entetes if titre in titres}
    for valeurs in lignes:
        if not any(v is not None for v in valeurs):
            continue
        yield {
            titre: valeurs[positions[titre]]
            if titre in positions and positions[titre] < len(valeurs)
            else None
            for titre in entetes
        }


@dataclass
class LectureBrute:
    date: dt.date
    index: Decimal
    commentaire: str = ""


@dataclass
class Bilan:
    maisons: int = 0
    compteurs: int = 0
    releves_crees: int = 0
    releves_maj: int = 0
    evenements: int = 0
    tarifs: int = 0
    avertissements: list[str] = field(default_factory=list)


class Command(BaseCommand):
    help = "Importe les relevés, événements et tarifs depuis compteur.xlsx."

    def add_arguments(self, parser):
        parser.add_argument(
            "fichier",
            nargs="?",
            default=str(settings.BASE_DIR / "compteur.xlsx"),
            help="Chemin du classeur (par défaut : compteur.xlsx à la racine).",
        )
        parser.add_argument(
            "--purger",
            action="store_true",
            help="Supprime maisons, compteurs et relevés avant l'import.",
        )

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError(
                "openpyxl est requis : .venv\\Scripts\\pip install openpyxl"
            ) from exc

        chemin = Path(options["fichier"])
        if not chemin.exists():
            raise CommandError(f"Fichier introuvable : {chemin}")

        self.bavard = options.get("verbosity", 1) >= 1
        classeur = load_workbook(chemin, data_only=True, read_only=True)
        bilan = Bilan()

        with transaction.atomic():
            if options["purger"]:
                Releve.objects.all().delete()
                Compteur.objects.all().delete()
                Evenement.objects.all().delete()
                Tarif.objects.all().delete()
                Maison.objects.all().delete()
                if self.bavard:
                    self.stdout.write(
                        self.style.WARNING("Données existantes supprimées.")
                    )

            if _est_un_export(classeur):
                self._importer_export(classeur, bilan)
            else:
                station = self._station()
                maisons = self._maisons(station, bilan)
                self._importer_releves(classeur, maisons, bilan)
                self._importer_tarifs(classeur, maisons, bilan)

        classeur.close()
        self._afficher(bilan)

    # -- création du référentiel -------------------------------------------

    def _station(self) -> StationMeteo:
        station, _ = StationMeteo.objects.get_or_create(
            nom=STATION_DEFAUT["nom"],
            defaults={
                "latitude": STATION_DEFAUT["latitude"],
                "longitude": STATION_DEFAUT["longitude"],
                "base_dj": Decimal(str(settings.BASE_DEGRES_JOURS)),
            },
        )
        return station

    def _maisons(self, station: StationMeteo, bilan: Bilan) -> dict[str, Maison]:
        resultat: dict[str, Maison] = {}
        for slug, donnees in MAISONS.items():
            maison, cree = Maison.objects.update_or_create(
                slug=slug,
                defaults={
                    "nom": donnees["nom"],
                    "nb_facades": donnees["nb_facades"],
                    "date_entree": donnees["date_entree"],
                    "date_sortie": donnees["date_sortie"],
                    "notes": donnees["notes"],
                    "station": station,
                },
            )
            resultat[slug] = maison
            bilan.maisons += int(cree)
        return resultat

    # -- relevés ------------------------------------------------------------

    def _importer_releves(self, classeur, maisons, bilan: Bilan) -> None:
        for bloc in BLOCS:
            if bloc.feuille not in classeur.sheetnames:
                bilan.avertissements.append(f"Feuille absente : {bloc.feuille}")
                continue
            feuille = classeur[bloc.feuille]
            maison = maisons[bloc.maison]

            for plage, colonne in bloc.colonnes.items():
                lectures = self._lire_bloc(feuille, bloc, colonne)
                if not lectures:
                    bilan.avertissements.append(
                        f"Aucun relevé lu : {bloc.feuille} / {bloc.maison} / {plage}"
                    )
                    continue
                self._enregistrer(maison, bloc, plage, lectures, bilan)

            self._extraire_evenements(feuille, bloc, maison, bilan)

    def _lire_bloc(self, feuille, bloc: Bloc, colonne: str) -> list[LectureBrute]:
        lectures: list[LectureBrute] = []
        for ligne in range(bloc.ligne_debut, bloc.ligne_fin + 1):
            jour = _vers_date(feuille[f"{bloc.col_date}{ligne}"].value)
            if jour is None:
                continue  # ligne de calcul intercalée
            index = _vers_decimal(feuille[f"{colonne}{ligne}"].value)
            if index is None:
                continue  # cellule « / » ou vide
            commentaires = [
                str(feuille[f"{col}{ligne}"].value).strip()
                for col in bloc.col_commentaires
                if isinstance(feuille[f"{col}{ligne}"].value, str)
            ]
            lectures.append(
                LectureBrute(
                    date=jour,
                    index=index,
                    commentaire=" / ".join(c for c in commentaires if c)[:200],
                )
            )
        lectures.sort(key=lambda l: l.date)
        return lectures

    @staticmethod
    def _decouper(lectures: list[LectureBrute]) -> list[list[LectureBrute]]:
        """Sépare la série à chaque recul d'index (remplacement de compteur)."""
        series: list[list[LectureBrute]] = [[]]
        for lecture in lectures:
            courante = series[-1]
            if courante and lecture.index < courante[-1].index:
                series.append([])
                courante = series[-1]
            courante.append(lecture)
        return [s for s in series if s]

    def _enregistrer(
        self,
        maison: Maison,
        bloc: Bloc,
        plage: str,
        lectures: list[LectureBrute],
        bilan: Bilan,
    ) -> None:
        coef = (
            Decimal(str(settings.COEF_CONVERSION_GAZ_KWH))
            if bloc.energie == Energie.GAZ
            else Decimal("1")
        )
        precedent: Compteur | None = None

        for rang, serie in enumerate(self._decouper(lectures), start=1):
            libelle = "" if rang == 1 else f"compteur {rang}"
            compteur, cree = Compteur.objects.update_or_create(
                maison=maison,
                energie=bloc.energie,
                plage=plage,
                libelle=libelle,
                defaults={
                    "unite": bloc.unite,
                    "coef_kwh": coef,
                    "date_pose": serie[0].date,
                    "date_depose": None,
                    "remplace": precedent,
                },
            )
            bilan.compteurs += int(cree)

            if precedent is not None:
                precedent.date_depose = serie[0].date
                precedent.save(update_fields=["date_depose"])

            for lecture in serie:
                source = (
                    SourceReleve.FOURNISSEUR
                    if "index" in lecture.commentaire.lower()
                    else SourceReleve.IMPORT
                )
                _, cree_releve = Releve.objects.update_or_create(
                    compteur=compteur,
                    date=lecture.date,
                    defaults={
                        "index": lecture.index,
                        "source": source,
                        "commentaire": lecture.commentaire,
                    },
                )
                if cree_releve:
                    bilan.releves_crees += 1
                else:
                    bilan.releves_maj += 1

            precedent = compteur

    def _extraire_evenements(self, feuille, bloc: Bloc, maison: Maison, bilan: Bilan):
        """Crée un événement pour chaque annotation évoquant des travaux."""
        for ligne in range(bloc.ligne_debut, bloc.ligne_fin + 1):
            jour = _vers_date(feuille[f"{bloc.col_date}{ligne}"].value)
            if jour is None:
                continue
            for colonne in bloc.col_commentaires:
                valeur = feuille[f"{colonne}{ligne}"].value
                if not isinstance(valeur, str):
                    continue
                texte = valeur.strip()
                if not any(mot in texte.lower() for mot in MOTS_CLES_EVENEMENT):
                    continue
                _, cree = Evenement.objects.get_or_create(
                    maison=maison,
                    date=jour,
                    libelle=texte[:150],
                    defaults={
                        "energie": bloc.energie,
                        # L'annotation n'est pas toujours portée par la ligne du
                        # jour concerné : certaines servent d'en-tête à une
                        # colonne de calcul. La date est donc à confirmer.
                        "description": (
                            f"Déduit automatiquement de la cellule {colonne}{ligne} "
                            f"de la feuille « {bloc.feuille} ». "
                            "Vérifiez la date : le classeur ne la porte pas toujours "
                            "sur la ligne du relevé correspondant."
                        ),
                    },
                )
                bilan.evenements += int(cree)

    # -- tarifs -------------------------------------------------------------

    def _importer_tarifs(self, classeur, maisons, bilan: Bilan) -> None:
        nom = "Prix par année"
        if nom not in classeur.sheetnames:
            bilan.avertissements.append(f"Feuille absente : {nom}")
            return
        feuille = classeur[nom]
        maison = maisons["collectivite"]

        for ligne_annee, ligne_montants, ligne_fournisseurs in BLOCS_TARIFS:
            valeur = feuille[f"A{ligne_annee}"].value
            if not isinstance(valeur, (int, float)) or not 1990 <= int(valeur) <= 2100:
                continue
            annee = int(valeur)

            for colonne, energie in COLONNES_TARIFS.items():
                montant = _vers_decimal(feuille[f"{colonne}{ligne_montants}"].value)
                if montant is None:
                    continue
                fournisseur = ""
                if ligne_fournisseurs:
                    brut = feuille[f"{colonne}{ligne_fournisseurs}"].value
                    if isinstance(brut, str):
                        fournisseur = brut.strip()

                _, cree = Tarif.objects.update_or_create(
                    maison=maison,
                    energie=energie,
                    date_debut=dt.date(annee, 1, 1),
                    defaults={
                        "date_fin": dt.date(annee, 12, 31),
                        "fournisseur": fournisseur,
                        "prix_unitaire": Decimal("0"),
                        "abonnement_mensuel": montant,
                        "notes": (
                            "Acompte mensuel repris du classeur. "
                            "Renseignez le prix unitaire pour valoriser les consommations."
                        ),
                    },
                )
                bilan.tarifs += int(cree)

        bilan.avertissements.append(
            "Les tarifs importés sont des acomptes mensuels : le prix unitaire "
            "(€/m³, €/kWh) reste à saisir dans l'admin pour chiffrer les consommations."
        )

    # -- classeur produit par l'export --------------------------------------

    def _importer_export(self, classeur, bilan: Bilan) -> None:
        """Relit un classeur produit par l'export de l'application.

        Contrairement à ``compteur.xlsx``, tout y est nommé : maisons et
        compteurs portent un code, et chaque bloc de relevés commence par celui
        de son compteur. Rien n'est donc deviné.
        """
        stations: dict[str, StationMeteo] = {}
        maisons: dict[str, Maison] = {}
        for ligne in _lignes(classeur[ex.FEUILLE_MAISONS], ex.ENTETES_MAISONS):
            code = _texte(ligne["Code"])
            if not code:
                continue
            station = None
            nom_station = _texte(ligne["Station météo"])
            if nom_station:
                station = stations.get(nom_station)
                if station is None:
                    station, _ = StationMeteo.objects.update_or_create(
                        nom=nom_station,
                        defaults={
                            "latitude": _vers_coordonnee(ligne["Latitude"]),
                            "longitude": _vers_coordonnee(ligne["Longitude"]),
                            "base_dj": _vers_decimal(ligne["Base des degrés-jours"])
                            or Decimal(str(settings.BASE_DEGRES_JOURS)),
                        },
                    )
                    stations[nom_station] = station
            maison, cree = Maison.objects.update_or_create(
                slug=code,
                defaults={
                    "nom": _texte(ligne["Nom"]) or code,
                    "adresse": _texte(ligne["Adresse"]),
                    "nb_facades": _vers_entier(ligne["Façades"]),
                    "surface_m2": _vers_entier(ligne["Surface chauffée (m²)"]),
                    "date_entree": _vers_jour(ligne["Occupée depuis"]),
                    "date_sortie": _vers_jour(ligne["Occupée jusqu'au"]),
                    "station": station,
                    "notes": _texte(ligne["Notes"]),
                },
            )
            maisons[code] = maison
            bilan.maisons += int(cree)

        compteurs: dict[str, Compteur] = {}
        remplacements: list[tuple[Compteur, str]] = []
        for ligne in _lignes(classeur[ex.FEUILLE_COMPTEURS], ex.ENTETES_COMPTEURS):
            code = _texte(ligne["Code"])
            maison = maisons.get(_texte(ligne["Maison"]))
            energie = _choix(Energie, ligne["Énergie"])
            plage = _choix(Plage, ligne["Plage"]) or Plage.UNIQUE
            if not code or maison is None or energie is None:
                bilan.avertissements.append(f"Compteur ignoré : {code or '(sans code)'}")
                continue
            compteur, cree = Compteur.objects.update_or_create(
                maison=maison,
                energie=energie,
                plage=plage,
                libelle=_texte(ligne["Libellé"]),
                defaults={
                    "numero": _texte(ligne["Numéro"]),
                    "unite": _texte(ligne["Unité"]) or "m³",
                    "coef_kwh": _vers_decimal(ligne["Coefficient kWh"]) or Decimal("1"),
                    "date_pose": _vers_jour(ligne["Posé le"]),
                    "date_depose": _vers_jour(ligne["Déposé le"]),
                },
            )
            compteurs[code] = compteur
            bilan.compteurs += int(cree)
            if _texte(ligne["Remplace"]):
                remplacements.append((compteur, _texte(ligne["Remplace"])))
        for compteur, code in remplacements:
            if code in compteurs:
                compteur.remplace = compteurs[code]
                compteur.save(update_fields=["remplace"])

        for _, titre in ex.FEUILLES_ENERGIE:
            if titre in classeur.sheetnames:
                self._relire_blocs(classeur[titre], compteurs, bilan)

        if ex.FEUILLE_EVENEMENTS in classeur.sheetnames:
            for ligne in _lignes(classeur[ex.FEUILLE_EVENEMENTS], ex.ENTETES_EVENEMENTS):
                maison = maisons.get(_texte(ligne["Maison"]))
                jour = _vers_jour(ligne["Date"])
                libelle = _texte(ligne["Libellé"])
                if maison is None or jour is None or not libelle:
                    continue
                _, cree = Evenement.objects.update_or_create(
                    maison=maison,
                    date=jour,
                    libelle=libelle[:150],
                    defaults={
                        "energie": _choix(Energie, ligne["Énergie"]) or "",
                        "description": _texte(ligne["Description"]),
                    },
                )
                bilan.evenements += int(cree)

        if ex.FEUILLE_TARIFS in classeur.sheetnames:
            for ligne in _lignes(classeur[ex.FEUILLE_TARIFS], ex.ENTETES_TARIFS):
                energie = _choix(Energie, ligne["Énergie"])
                debut = _vers_jour(ligne["Début"])
                if energie is None or debut is None:
                    continue
                _, cree = Tarif.objects.update_or_create(
                    maison=maisons.get(_texte(ligne["Maison"])),
                    energie=energie,
                    date_debut=debut,
                    defaults={
                        "date_fin": _vers_jour(ligne["Fin"]),
                        "fournisseur": _texte(ligne["Fournisseur"]),
                        "prix_unitaire": _vers_decimal(ligne["Prix unitaire (€)"], 5)
                        or Decimal("0"),
                        "abonnement_mensuel": _vers_decimal(ligne["Abonnement (€/mois)"], 2)
                        or Decimal("0"),
                        "notes": _texte(ligne["Notes"]),
                    },
                )
                bilan.tarifs += int(cree)

        if ex.FEUILLE_DJ in classeur.sheetnames:
            self._relire_degres_jours(classeur[ex.FEUILLE_DJ], stations, bilan)

    def _relire_blocs(self, feuille, compteurs: dict[str, Compteur], bilan: Bilan) -> None:
        """Parcourt les blocs d'une feuille d'énergie, chacun ouvert par un code."""
        sources = {libelle: code for code, libelle in SourceReleve.choices}
        compteur: Compteur | None = None
        for valeurs in feuille.iter_rows(values_only=True):
            premiere = valeurs[0] if valeurs else None
            if isinstance(premiere, str) and premiere.strip() in compteurs:
                compteur = compteurs[premiere.strip()]
                continue
            jour = _vers_jour(premiere)
            if compteur is None or jour is None:
                continue  # ligne d'en-tête, ligne vide ou bloc inconnu
            index = _vers_decimal(valeurs[1] if len(valeurs) > 1 else None)
            if index is None:
                continue
            colonne = lambda rang: valeurs[rang] if len(valeurs) > rang else None  # noqa: E731
            _, cree = Releve.objects.update_or_create(
                compteur=compteur,
                date=jour,
                defaults={
                    "index": index,
                    "annuel": _texte(colonne(6)).lower() == "oui",
                    "source": sources.get(_texte(colonne(7)), SourceReleve.IMPORT),
                    "commentaire": _texte(colonne(8))[:200],
                },
            )
            if cree:
                bilan.releves_crees += 1
            else:
                bilan.releves_maj += 1

    def _relire_degres_jours(self, feuille, stations: dict[str, StationMeteo], bilan: Bilan):
        """Les degrés-jours se retéléchargent, mais les relire évite d'attendre."""
        connus: dict[int, set[dt.date]] = {}
        nouveaux: list[DegreJour] = []
        for ligne in _lignes(feuille, ex.ENTETES_DJ):
            station = stations.get(_texte(ligne["Station"]))
            jour = _vers_jour(ligne["Date"])
            valeur = _vers_decimal(ligne["Degrés-jours"])
            if station is None or jour is None or valeur is None:
                continue
            if station.pk not in connus:
                connus[station.pk] = set(
                    DegreJour.objects.filter(station=station).values_list("date", flat=True)
                )
            if jour in connus[station.pk]:
                continue
            temperature = ligne["Température moyenne (°C)"]
            nouveaux.append(
                DegreJour(
                    station=station,
                    date=jour,
                    temperature_moyenne=None
                    if temperature is None
                    else Decimal(f"{float(temperature):.2f}"),
                    dj=valeur,
                )
            )
            connus[station.pk].add(jour)
        DegreJour.objects.bulk_create(nouveaux, batch_size=1000)
        if nouveaux:
            bilan.avertissements.append(f"{len(nouveaux)} degrés-jours repris du classeur.")

    # -- sortie -------------------------------------------------------------

    def _afficher(self, bilan: Bilan) -> None:
        if not self.bavard:
            return
        self.stdout.write(self.style.SUCCESS("Import terminé."))
        self.stdout.write(f"  maisons créées      : {bilan.maisons}")
        self.stdout.write(f"  compteurs créés     : {bilan.compteurs}")
        self.stdout.write(f"  relevés créés       : {bilan.releves_crees}")
        self.stdout.write(f"  relevés mis à jour  : {bilan.releves_maj}")
        self.stdout.write(f"  événements          : {bilan.evenements}")
        self.stdout.write(f"  tarifs              : {bilan.tarifs}")
        for message in bilan.avertissements:
            self.stdout.write(self.style.WARNING(f"  ! {message}"))
        self.stdout.write(
            "\nÉtape suivante : python manage.py sync_degres_jours "
            "(indispensable pour la normalisation climatique)."
        )
