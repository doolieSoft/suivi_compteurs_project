"""Tests de l'export Excel et de sa réimportation.

La promesse de l'export est d'être une sauvegarde : exporter, tout effacer,
puis réimporter doit rendre exactement les mêmes données.
"""
from __future__ import annotations

import datetime as dt
import io
import tempfile
from decimal import Decimal
from pathlib import Path

from django.core.management import call_command
from django.test import TestCase
from openpyxl import load_workbook

from suivi.models import (
    Compteur,
    DegreJour,
    Energie,
    Evenement,
    Maison,
    Plage,
    Releve,
    SourceReleve,
    Tarif,
)
from suivi.services import export_excel
from suivi.tests import fabrique as f


def _instantane() -> dict:
    """Photographie comparable des données, indépendante des identifiants."""
    return {
        "maisons": sorted(
            Maison.objects.values_list(
                "slug", "nom", "nb_facades", "date_entree", "date_sortie", "station__nom"
            )
        ),
        "compteurs": sorted(
            Compteur.objects.values_list(
                "maison__slug", "energie", "plage", "libelle", "unite", "coef_kwh",
                "date_pose", "date_depose", "remplace__libelle",
            )
        ),
        "releves": sorted(
            Releve.objects.values_list(
                "compteur__maison__slug", "compteur__energie", "compteur__plage",
                "compteur__libelle", "date", "index", "annuel", "source", "commentaire",
            )
        ),
        "evenements": sorted(
            Evenement.objects.values_list("maison__slug", "date", "energie", "libelle")
        ),
        "tarifs": sorted(
            Tarif.objects.values_list(
                "maison__slug", "energie", "date_debut", "date_fin", "prix_unitaire",
                "abonnement_mensuel",
            ),
            key=repr,  # le tarif commun à toutes les maisons n'a pas de maison
        ),
        "dj": DegreJour.objects.count(),
    }


class ExportReimportTest(TestCase):
    def setUp(self):
        station = f.station()
        f.climat(station, dt.date(2023, 1, 1), dt.date(2023, 3, 31))
        maison = f.maison(station, nom="Maison d'essai")
        ancienne = f.maison(station, nom="Ancienne")
        ancienne.date_sortie = dt.date(2022, 10, 31)
        ancienne.save()

        gaz = f.compteur(maison, Energie.GAZ)
        gaz.date_depose = dt.date(2023, 2, 1)
        gaz.save()
        gaz2 = f.compteur(maison, Energie.GAZ, libelle="compteur 2")
        gaz2.remplace = gaz
        gaz2.save()
        haut = f.compteur(ancienne, Energie.ELECTRICITE, plage=Plage.HAUT, unite="kWh")
        bas = f.compteur(ancienne, Energie.ELECTRICITE, plage=Plage.BAS, unite="kWh")

        for compteur, valeurs in (
            (gaz, [(dt.date(2023, 1, 3), "1000.125"), (dt.date(2023, 1, 31), "1100.5")]),
            (gaz2, [(dt.date(2023, 2, 1), "0"), (dt.date(2023, 3, 1), "80.001")]),
            (haut, [(dt.date(2021, 12, 31), "500.1"), (dt.date(2022, 12, 31), "900.4")]),
            (bas, [(dt.date(2021, 12, 31), "300"), (dt.date(2022, 12, 31), "650.7")]),
        ):
            for jour, index in valeurs:
                Releve.objects.create(compteur=compteur, date=jour, index=Decimal(index))
        Releve.objects.filter(compteur=gaz, date=dt.date(2023, 1, 3)).update(
            annuel=True, source=SourceReleve.FOURNISSEUR, commentaire="index envoyé"
        )
        Evenement.objects.create(
            maison=maison, date=dt.date(2023, 2, 1), energie=Energie.GAZ,
            libelle="Remplacement du compteur",
        )
        Tarif.objects.create(
            maison=maison, energie=Energie.GAZ, date_debut=dt.date(2023, 1, 1),
            prix_unitaire=Decimal("1.23456"), abonnement_mensuel=Decimal("12.50"),
        )
        Tarif.objects.create(energie=Energie.EAU, date_debut=dt.date(2023, 1, 1))

    def _exporter(self) -> Path:
        dossier = Path(tempfile.mkdtemp())
        chemin = dossier / "export.xlsx"
        chemin.write_bytes(export_excel.classeur())
        return chemin

    def test_exporter_effacer_reimporter_rend_les_memes_donnees(self):
        avant = _instantane()
        chemin = self._exporter()

        call_command("import_excel", str(chemin), purger=True, verbosity=0)
        DegreJour.objects.all().delete()
        call_command("import_excel", str(chemin), verbosity=0)

        self.assertEqual(_instantane(), avant)

    def test_reimporter_deux_fois_ne_duplique_rien(self):
        avant = _instantane()
        chemin = self._exporter()
        call_command("import_excel", str(chemin), verbosity=0)
        call_command("import_excel", str(chemin), verbosity=0)
        self.assertEqual(_instantane(), avant)

    def test_une_ligne_ajoutee_a_la_main_est_reprise(self):
        chemin = self._exporter()
        classeur = load_workbook(chemin)
        feuille = classeur["Gaz"]
        # Fin du bloc du second compteur de gaz : on ajoute un relevé dessous.
        derniere = max(
            c.row for (c,) in feuille.iter_rows(min_col=1, max_col=1) if c.value is not None
        )
        feuille.insert_rows(derniere + 1)
        feuille.cell(derniere + 1, 1, dt.datetime(2023, 3, 20))
        feuille.cell(derniere + 1, 2, 95.5)
        classeur.save(chemin)

        call_command("import_excel", str(chemin), verbosity=0)
        releve = Releve.objects.get(date=dt.date(2023, 3, 20))
        self.assertEqual(releve.compteur.libelle, "compteur 2")
        self.assertEqual(releve.index, Decimal("95.500"))

    def test_la_presentation_reprend_celle_du_classeur_d_origine(self):
        classeur = load_workbook(io.BytesIO(export_excel.classeur()))
        self.assertEqual(classeur.sheetnames[:3], ["Eau", "Gaz", "Électricité"])
        feuille = classeur["Gaz"]
        self.assertEqual(feuille["A2"].value, "Date")
        self.assertEqual(feuille["B2"].value, "Index")
        # Les écarts sont des formules, comme dans compteur.xlsx.
        self.assertEqual(feuille["C4"].value, "=B4-B3")
